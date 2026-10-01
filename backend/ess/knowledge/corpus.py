"""Corpus documents for business-identity mining.

A corpus mixes what the company wrote (own quotations, company documents, sent mail), what its
customers wrote (inbound mail) and optional online context. ``SOURCE_WEIGHTS`` encodes how much
each kind of source can prove: the company's own quotations decide, web pages are context only.
"""
from __future__ import annotations

import email
import email.policy
import logging
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

from ess.knowledge.base import detect_language
from ess.knowledge.text import domain_matches, email_domain, quoted_cut

log = logging.getLogger(__name__)

SOURCE_TYPES = ("own_quotation", "company_doc", "sent_email", "inbound_email", "web")
SOURCE_WEIGHTS: dict[str, float] = {
    "own_quotation": 1.0,
    "company_doc": 0.9,
    "sent_email": 0.8,
    "inbound_email": 0.5,
    "web": 0.2,
}
#: What the company itself wrote.
OWN_SOURCE_TYPES = frozenset({"own_quotation", "company_doc", "sent_email"})
#: Sources that show work actually offered/delivered (not just claimed).
DELIVERED_SOURCE_TYPES = frozenset({"own_quotation", "sent_email"})
#: What customers wrote.
CUSTOMER_SOURCE_TYPES = frozenset({"inbound_email"})
EMAIL_SOURCE_TYPES = frozenset({"sent_email", "inbound_email"})

DOCUMENT_EXTENSIONS = frozenset(
    {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".xlsm", ".pptx", ".txt", ".md", ".markdown", ".csv", ".html",
     ".htm", ".eml", ".msg", ".rtf", ".odt"}
)
FALLBACK_EXTENSIONS = frozenset({".txt", ".md", ".markdown", ".csv", ".docx", ".pdf", ".html", ".htm", ".eml"})


@dataclass
class CorpusDoc:
    """One text the miner can learn from."""

    source_type: str  # own_quotation | company_doc | sent_email | inbound_email | web
    source_id: str  # message id, file path relative to the scanned folder, or URL
    label: str  # human label shown next to evidence ("Quotation NW/26/0101.pdf", "RFQ: ...")
    text: str
    date: str | None = None  # ISO-8601
    language: str | None = None  # detected when not given
    meta: dict[str, Any] = field(default_factory=dict)  # from_email, to, thread_id, path, url...

    def __post_init__(self) -> None:
        if self.source_type not in SOURCE_WEIGHTS:
            raise ValueError(f"unknown source_type {self.source_type!r}; expected one of {SOURCE_TYPES}")
        self.source_id = str(self.source_id)
        self.text = "" if self.text is None else str(self.text)
        self.label = str(self.label or self.source_id)
        if isinstance(self.date, datetime):
            dt = self.date if self.date.tzinfo else self.date.replace(tzinfo=timezone.utc)
            self.date = dt.isoformat()
        elif isinstance(self.date, date):
            self.date = self.date.isoformat()
        if self.language is None:
            self.language = detect_language(self.text)
        if self.meta is None:
            self.meta = {}

    @property
    def weight(self) -> float:
        return SOURCE_WEIGHTS[self.source_type]

    @property
    def is_own(self) -> bool:
        return self.source_type in OWN_SOURCE_TYPES

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


# --------------------------------------------------------------------------------------------
# Folders of company files (old quotations, profiles, manuals)
# --------------------------------------------------------------------------------------------

_QUOTATION_NAME_RE = re.compile(r"(?i)quot|qtn|offer|proposal|عرض|تسعير")
_QUOTATION_TEXT_RE = re.compile(
    r"(?i)\b(?:quotation|our\s+offer|we\s+are\s+pleased\s+to\s+(?:quote|offer|submit)|price\s+offer|commercial\s+offer)\b"
    r"|عرض\s+سعر|عرض\s+أسعار|коммерческое\s+предложение")


@lru_cache(maxsize=1)
def _documents_extractor() -> Callable[[Path], Any] | None:
    try:
        from ess.documents.extract import extract_document  # type: ignore
    except Exception:  # package not available yet: use the small built-in readers
        return None
    return extract_document


def _text_of_extract(result: Any) -> str:
    if result is None:
        return ""
    get = result.get if isinstance(result, dict) else (lambda k, d=None: getattr(result, k, d))
    text = get("text") or ""
    if not text:
        pages = get("pages") or []
        parts = []
        for p in pages:
            t = p.get("text") if isinstance(p, dict) else getattr(p, "text", "")
            if t:
                parts.append(str(t))
        text = "\n\n".join(parts)
    return str(text or "")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _html_text(markup: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(markup, "lxml")
    for tag in soup(["script", "style", "noscript", "template", "svg"]):
        tag.decompose()
    text = soup.get_text("\n")
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _read_fallback(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".txt", ".md", ".markdown", ".csv"):
        return _decode(path.read_bytes())
    if ext in (".html", ".htm"):
        return _html_text(_decode(path.read_bytes()))
    if ext == ".docx":
        import docx  # python-docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                cells = []
                for c in row.cells:
                    t = c.text.strip()
                    if t and (not cells or cells[-1] != t):
                        cells.append(t)
                if cells:
                    parts.append(" | ".join(cells))
        for section in d.sections:
            for hf in (section.header, section.footer):
                try:
                    parts.extend(p.text for p in hf.paragraphs)
                except Exception:
                    pass
        return "\n".join(p for p in parts if p is not None)
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".eml":
        msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
        body = msg.get_body(preferencelist=("plain", "html"))
        content = body.get_content() if body is not None else ""
        if body is not None and body.get_content_type() == "text/html":
            content = _html_text(content)
        subject = str(msg.get("subject") or "")
        return f"{subject}\n{content}" if subject else content
    raise ValueError(f"unsupported file type {ext}")


def read_document_text(path: Path | str) -> str:
    """Text of one file: ``ess.documents.extract.extract_document`` when available, else built-in
    readers for .txt/.md/.csv/.html/.docx/.pdf/.eml."""
    p = Path(path)
    extractor = _documents_extractor()
    text = ""
    if extractor is not None:
        try:
            text = _text_of_extract(extractor(p))
        except Exception as exc:  # fall through to the built-in readers
            log.info("extract_document failed for %s: %s", p, exc)
    if not text.strip() and p.suffix.lower() in FALLBACK_EXTENSIONS:
        text = _read_fallback(p)
    return text


def guess_source_type(path: Path | str, text: str) -> str:
    """``own_quotation`` for files that look like the company's quotations, else ``company_doc``."""
    name = Path(path).name
    head = (text or "")[:3000]
    if _QUOTATION_NAME_RE.search(name) or _QUOTATION_TEXT_RE.search(head):
        return "own_quotation"
    return "company_doc"


def _iter_files(root: Path, recursive: bool) -> Iterator[Path]:
    it = root.rglob("*") if recursive else root.glob("*")
    for p in sorted(it):
        if not p.is_file():
            continue
        if any(part.startswith(".") for part in p.relative_to(root).parts) or p.name.startswith("~$"):
            continue
        yield p


def build_corpus_from_folder(path: Path | str, source_type: str = "company_doc", limit: int | None = None, *,
                             recursive: bool = True, max_chars: int = 400_000,
                             extensions: Iterable[str] | None = None) -> list[CorpusDoc]:
    """Read company files (old quotations, profiles, manuals) into corpus documents.

    ``source_type="auto"`` tells quotations (file name or wording) from other company documents.
    Unreadable files are skipped and logged; files are visited in a stable (sorted) order.
    """
    root = Path(path).expanduser()
    if source_type != "auto" and source_type not in SOURCE_WEIGHTS:
        raise ValueError(f"unknown source_type {source_type!r}")
    if root.is_file():
        files: Iterable[Path] = [root]
        base = root.parent
    elif root.is_dir():
        files = _iter_files(root, recursive)
        base = root
    else:
        raise FileNotFoundError(str(root))
    allowed = {e.lower() if e.startswith(".") else "." + e.lower() for e in extensions} if extensions else None
    docs: list[CorpusDoc] = []
    for f in files:
        if limit is not None and len(docs) >= limit:
            break
        ext = f.suffix.lower()
        if allowed is not None:
            if ext not in allowed:
                continue
        elif ext not in DOCUMENT_EXTENSIONS:
            continue
        try:
            text = read_document_text(f)
        except Exception as exc:
            log.warning("skipping %s: %s", f, exc)
            continue
        text = (text or "").strip()
        if not text:
            continue
        text = text[:max_chars]
        stype = guess_source_type(f, text) if source_type == "auto" else source_type
        try:
            mtime = datetime.fromtimestamp(f.stat().st_mtime, tz=timezone.utc).date().isoformat()
        except OSError:
            mtime = None
        rel = f.relative_to(base).as_posix() if f.is_relative_to(base) else f.name
        docs.append(CorpusDoc(source_type=stype, source_id=rel, label=f.name, text=text, date=mtime,
                              meta={"path": str(f)}))
    return docs


# --------------------------------------------------------------------------------------------
# Mailbox messages
# --------------------------------------------------------------------------------------------

_REPLY_SUBJECT_RE = re.compile(r"^\s*(?:re|aw|sv|antw|رد)\s*:", re.I)
_FWD_SUBJECT_RE = re.compile(r"^\s*(?:fw|fwd|wg|tr|rv|إعادة توجيه)\s*:", re.I)
_ADDR_RE = re.compile(r"[\w.+'-]+@[\w-]+(?:\.[\w-]+)+")


def _field(obj: Any, *names: str) -> Any:
    for n in names:
        if isinstance(obj, dict):
            v = obj.get(n)
        else:
            v = getattr(obj, n, None)
        if v not in (None, "", [], {}):
            return v
    return None


def _addr(value: Any) -> str:
    if not value:
        return ""
    m = _ADDR_RE.search(str(value))
    return m.group(0).lower() if m else ""


def _iso(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return (value if value.tzinfo else value.replace(tzinfo=timezone.utc)).isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def from_mail_messages(messages: Iterable[Any], own_domains: Iterable[str], *, account: str | None = None,
                       include_subject: bool = True, max_chars: int = 20_000) -> list[CorpusDoc]:
    """Turn mailbox messages (``ess.sources.base.MailMessage`` or snapshot e-mail dicts) into docs.

    Outbound mail (``direction="outbound"``, label ``SENT``, or a sender on an own domain) becomes
    ``sent_email`` and keeps ONLY what the company wrote (quoted history and forwarded content
    are cut). Inbound mail becomes ``inbound_email`` without the quoted reply history (forwarded
    RFQs are kept: they are customer-side text). Reply subjects are not included (they repeat
    someone else's words). Empty messages are skipped.
    """
    own = [d.lower().strip().lstrip("@") for d in (own_domains or []) if d and str(d).strip()]
    acct = (account or "").lower()
    out: list[CorpusDoc] = []
    for m in messages or []:
        mid = _field(m, "id", "message_id", "source_id")
        if mid is None:
            continue
        sender = _addr(_field(m, "from_email", "sender", "from", "from_address"))
        labels = {str(x).upper() for x in (_field(m, "labels") or [])}
        direction = str(_field(m, "direction") or "").lower()
        # "inbound" is also the model default, so an own-domain sender or the SENT label still wins
        outbound = (direction == "outbound" or "SENT" in labels or (bool(acct) and sender == acct)
                    or domain_matches(email_domain(sender), own))
        direction = "outbound" if outbound else "inbound"
        body = _field(m, "body_text", "body", "text") or ""
        if not body:
            html = _field(m, "body_html", "html")
            if html:
                try:
                    from ess.knowledge.corpus import _html_text as _h

                    body = _h(str(html))
                except Exception:
                    body = ""
        if not body:
            body = _field(m, "snippet") or ""
        body = str(body)
        subject = str(_field(m, "subject") or "")
        if direction == "outbound":
            stype = "sent_email"
            own_text = body[:quoted_cut(body, "all")].strip()
            skip_subject = bool(_REPLY_SUBJECT_RE.match(subject) or _FWD_SUBJECT_RE.match(subject))
        else:
            stype = "inbound_email"
            own_text = body[:quoted_cut(body, "reply")].strip()
            skip_subject = bool(_REPLY_SUBJECT_RE.match(subject))
        if not own_text and (skip_subject or not subject):
            continue
        text = own_text
        if include_subject and subject and not skip_subject:
            text = f"{subject}\n{own_text}" if own_text else subject
        to = _field(m, "to") or []
        if isinstance(to, str):
            to = [to]
        label = ("Sent: " if stype == "sent_email" else "Email: ") + (subject[:110] or str(mid))
        out.append(CorpusDoc(
            source_type=stype, source_id=str(mid), label=label, text=text[:max_chars], date=_iso(_field(m, "date")),
            meta={"from_email": sender, "from_domain": email_domain(sender), "to": [_addr(t) for t in to if _addr(t)],
                  "thread_id": _field(m, "thread_id"), "subject": subject, "direction": direction,
                  "view_url": _field(m, "view_url")}))
    return out
