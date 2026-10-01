"""Read every file a customer sends: text, pages, tables, BOQ items, images and metadata.

    doc = extract_document(path)
    doc.kind        # pdf | docx | spreadsheet | csv | image | archive | email | cad | text | presentation | unsupported
    doc.doc_kind    # classify_document(): drawing | boq | specification | tender_doc | addendum | photo | cad | correspondence | other
    doc.pages       # [{n (1-based), text, is_drawing_like, ...}]   (PDF pages, sheets, slides)
    doc.boq_items   # [{ref, description, qty, unit, rate, amount, section, row, relevant, match, ...}]
    png = render_pdf_page(path, page_number=1, dpi=110)

Archives and e-mails are unpacked safely (zip-slip guard, size caps, executables skipped) into
``<file>_unpacked/`` next to the file (or ``extract_dir``), and supported members are extracted
recursively into ``doc.children``. Content problems never raise: they become ``warnings``.
"""
from __future__ import annotations

import csv
import email
import email.policy
import io
import mimetypes
import os
import re
import stat
import zipfile
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import Any

from pydantic import BaseModel, Field

from ..browser.safety import blocked_reason, sanitize_filename, unique_path
from .boq import is_relevant, parse_boq_rows
from .cad import read_cad
from .pdf import extract_pdf, render_pdf_page

__all__ = ["DocExtract", "classify_document", "detect_kind", "extract_document", "render_pdf_page"]

MAX_TEXT_CHARS = 5_000_000
MAX_SHEET_ROWS = 20_000
MAX_SHEET_COLS = 60
MAX_ZIP_MEMBERS = 2_000
MAX_ZIP_MEMBER_BYTES = 1_000_000_000
MAX_ZIP_TOTAL_BYTES = 3_000_000_000
MAX_DEPTH = 3

_EXT_KIND = {
    ".pdf": "pdf",
    ".docx": "docx", ".docm": "docx", ".dotx": "docx",
    ".xlsx": "spreadsheet", ".xlsm": "spreadsheet", ".xltx": "spreadsheet", ".xls": "spreadsheet", ".xlsb": "spreadsheet",
    ".ods": "spreadsheet",
    ".csv": "csv", ".tsv": "csv",
    ".jpg": "image", ".jpeg": "image", ".png": "image", ".gif": "image", ".bmp": "image", ".tif": "image",
    ".tiff": "image", ".webp": "image", ".heic": "image", ".heif": "image",
    ".zip": "archive", ".rar": "archive", ".7z": "archive", ".tar": "archive", ".gz": "archive", ".tgz": "archive",
    ".eml": "email", ".msg": "email",
    ".dwg": "cad", ".dxf": "cad", ".rvt": "cad", ".rfa": "cad", ".ifc": "cad", ".dgn": "cad", ".skp": "cad",
    ".nwd": "cad", ".nwc": "cad", ".dwf": "cad", ".dwfx": "cad", ".step": "cad", ".stp": "cad", ".iges": "cad",
    ".igs": "cad",
    ".txt": "text", ".md": "text", ".rtf": "text", ".html": "text", ".htm": "text", ".xml": "text", ".json": "text",
    ".pptx": "presentation",
    ".doc": "unsupported", ".ppt": "unsupported",
}


class _DictAccess:
    def __getitem__(self, key: str) -> Any:
        try:
            return getattr(self, key)
        except AttributeError as exc:
            raise KeyError(key) from exc

    def get(self, key: str, default: Any = None) -> Any:
        return getattr(self, key, default)


class DocExtract(_DictAccess, BaseModel):
    path: str
    name: str
    kind: str
    doc_kind: str = "other"
    text: str = ""
    pages: list[dict] = Field(default_factory=list)
    tables: list[dict] = Field(default_factory=list)
    boq_items: list[dict] = Field(default_factory=list)
    images: list[dict] = Field(default_factory=list)
    meta: dict = Field(default_factory=dict)
    warnings: list[str] = Field(default_factory=list)
    children: list["DocExtract"] = Field(default_factory=list)

    @property
    def relevant_boq_items(self) -> list[dict]:
        return [b for b in self.boq_items if b.get("relevant")]

    @property
    def drawing_pages(self) -> list[int]:
        return [p["n"] for p in self.pages if p.get("is_drawing_like")]


# --------------------------------------------------------------------------- helpers
def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
            return int(value)
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _row_text(row: list[Any]) -> str:
    cells = [str(_json_value(v)).strip() for v in row if v not in (None, "")]
    return " | ".join(c for c in cells if c)


def detect_kind(path: Path) -> str:
    """File family from the extension, or from magic bytes for bare/odd names."""
    ext = path.suffix.lower()
    if ext in _EXT_KIND:
        return _EXT_KIND[ext]
    try:
        head = path.read_bytes()[:16] if path.stat().st_size < 64 else path.open("rb").read(16)
    except OSError:
        return "unsupported"
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as zf:
                names = set(zf.namelist())
            if "word/document.xml" in names:
                return "docx"
            if "xl/workbook.xml" in names:
                return "spreadsheet"
            if "ppt/presentation.xml" in names:
                return "presentation"
        except zipfile.BadZipFile:
            return "unsupported"
        return "archive"
    if head[:3] == b"\xff\xd8\xff" or head.startswith(b"\x89PNG") or head[:4] in (b"II*\x00", b"MM\x00*"):
        return "image"
    if head.startswith(b"AC10"):
        return "cad"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "email" if ext == ".msg" else "unsupported"
    return "unsupported"


def _child_dir(path: Path, extract_dir: Path | None) -> Path:
    if extract_dir is not None:
        return Path(extract_dir)
    return path.parent / f"{path.name}_unpacked"


# --------------------------------------------------------------------------- formats
def _extract_docx(path: Path, out: DocExtract) -> None:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines: list[str] = []
    table_no = 0

    def table_rows(table: Any) -> list[list[str]]:
        rows: list[list[str]] = []
        for row in table.rows:
            values: list[str] = []
            previous = None
            for cell in row.cells:
                text = re.sub(r"\s+", " ", cell.text).strip()
                values.append("" if cell._tc is previous else text)  # merged cells repeat
                previous = cell._tc
            rows.append(values)
        return rows

    try:
        blocks = list(document.iter_inner_content())
    except AttributeError:  # python-docx < 1.0
        blocks = []
        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                blocks.append(Paragraph(child, document))
            elif tag == "tbl":
                blocks.append(Table(child, document))
    for block in blocks:
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                lines.append(text)
        elif isinstance(block, Table):
            table_no += 1
            rows = table_rows(block)
            out.tables.append({"name": f"Table {table_no}", "rows": rows})
            lines += [" | ".join(c for c in r if c) for r in rows if any(r)]
            out.boq_items += parse_boq_rows(rows, sheet=f"Table {table_no}")
    extra: list[str] = []
    for section in document.sections:
        for part in (section.header, section.footer):
            try:
                for para in part.paragraphs:
                    if para.text.strip() and para.text.strip() not in extra:
                        extra.append(para.text.strip())
            except Exception:
                continue
    if extra:
        lines += ["", "[header/footer]"] + extra
    out.text = "\n".join(lines)
    out.pages = [{"n": 1, "text": out.text, "is_drawing_like": False}]
    props = document.core_properties
    out.meta.update(
        {
            "author": props.author or None,
            "title": props.title or None,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
            "tables": table_no,
        }
    )
    for rel in document.part.rels.values():
        if "image" in rel.reltype and not rel.is_external:
            part = rel.target_part
            out.images.append({"name": PurePosixPath(str(part.partname)).name, "mime": part.content_type,
                               "size": len(part.blob)})


def _extract_xlsx(path: Path, out: DocExtract) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheet_texts: list[str] = []
        for n, sheet in enumerate(workbook.worksheets, start=1):
            rows: list[list[Any]] = []
            numbers: list[int] = []
            truncated = False
            for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if r_idx > MAX_SHEET_ROWS:
                    truncated = True
                    break
                values = [_json_value(v) for v in (row or ())[:MAX_SHEET_COLS]]
                if any(v not in (None, "") for v in values):
                    while values and values[-1] in (None, ""):
                        values.pop()
                    rows.append(values)
                    numbers.append(r_idx)
            if truncated:
                out.warnings.append(f"sheet '{sheet.title}': only the first {MAX_SHEET_ROWS} rows were read")
            state = getattr(sheet, "sheet_state", "visible")
            out.tables.append({"name": sheet.title, "rows": rows, "row_numbers": numbers, "state": state})
            text = "\n".join(_row_text(r) for r in rows)
            sheet_texts.append(f"## {sheet.title}\n{text}")
            out.pages.append({"n": n, "text": text, "is_drawing_like": False, "label": sheet.title})
            out.boq_items += parse_boq_rows(rows, row_numbers=numbers, sheet=sheet.title)
        out.text = "\n\n".join(sheet_texts)
        props = workbook.properties
        out.meta.update(
            {
                "sheets": [ws.title for ws in workbook.worksheets],
                "author": getattr(props, "creator", None) or None,
                "title": getattr(props, "title", None) or None,
                "created": props.created.isoformat() if getattr(props, "created", None) else None,
            }
        )
    finally:
        workbook.close()


def _extract_xls(path: Path, out: DocExtract) -> None:
    try:
        import xlrd  # type: ignore
    except ImportError:
        out.warnings.append(".xls (Excel 97-2003) needs the 'xlrd' package; save it as .xlsx to read it")
        return
    book = xlrd.open_workbook(str(path))
    texts = []
    for n, sheet in enumerate(book.sheets(), start=1):
        rows = [[_json_value(sheet.cell_value(r, c)) for c in range(min(sheet.ncols, MAX_SHEET_COLS))]
                for r in range(min(sheet.nrows, MAX_SHEET_ROWS))]
        numbers = list(range(1, len(rows) + 1))
        out.tables.append({"name": sheet.name, "rows": rows, "row_numbers": numbers})
        text = "\n".join(_row_text(r) for r in rows if any(v not in (None, "") for v in r))
        texts.append(f"## {sheet.name}\n{text}")
        out.pages.append({"n": n, "text": text, "is_drawing_like": False, "label": sheet.name})
        out.boq_items += parse_boq_rows(rows, row_numbers=numbers, sheet=sheet.name)
    out.text = "\n\n".join(texts)
    out.meta["sheets"] = [s.name for s in book.sheets()]


def _decode_text_bytes(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    for encoding in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_csv(path: Path, out: DocExtract) -> None:
    text = _decode_text_bytes(path.read_bytes()[:50_000_000])
    sample = text[:8_192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= MAX_SHEET_ROWS:
            out.warnings.append(f"only the first {MAX_SHEET_ROWS} rows were read")
            break
        rows.append([c.strip() for c in row[:MAX_SHEET_COLS]])
    numbers = list(range(1, len(rows) + 1))
    out.tables.append({"name": path.name, "rows": rows, "row_numbers": numbers})
    out.text = "\n".join(_row_text(r) for r in rows if any(r))
    out.pages = [{"n": 1, "text": out.text, "is_drawing_like": False}]
    out.boq_items = parse_boq_rows(rows, row_numbers=numbers, sheet=path.name)


def _extract_image(path: Path, out: DocExtract) -> None:
    out.warnings.append("no OCR: text inside images is not extracted")
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        return
    try:
        with Image.open(str(path)) as image:
            info: dict[str, Any] = {"name": path.name, "width": image.width, "height": image.height,
                                    "format": image.format, "mode": image.mode}
            exif = image.getexif()
            if exif:
                camera = " ".join(str(exif.get(tag, "")).strip() for tag in (0x010F, 0x0110)).strip()
                if camera:
                    info["camera"] = camera
                try:
                    sub = exif.get_ifd(0x8769)
                    taken = sub.get(0x9003) or exif.get(0x0132)
                except Exception:
                    taken = exif.get(0x0132)
                if taken:
                    info["taken"] = str(taken)
                if 0x8825 in exif:
                    info["gps"] = True
            out.images.append(info)
            out.meta.update({k: v for k, v in info.items() if k != "name"})
            frames = getattr(image, "n_frames", 1)
            if frames and frames > 1:
                out.meta["frames"] = frames
    except Exception as exc:
        hint = " (HEIC needs the pillow-heif package)" if path.suffix.lower() in (".heic", ".heif") else ""
        out.warnings.append(f"image could not be opened{hint}: {type(exc).__name__}")


def _extract_text_file(path: Path, out: DocExtract) -> None:
    text = _decode_text_bytes(path.read_bytes()[:MAX_TEXT_CHARS])
    if path.suffix.lower() in (".html", ".htm"):
        from ..sources.base import html_to_text

        text = html_to_text(text)
    elif path.suffix.lower() == ".rtf":
        text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", text)
    out.text = text
    out.pages = [{"n": 1, "text": text, "is_drawing_like": False}]


def _extract_pptx(path: Path, out: DocExtract) -> None:
    with zipfile.ZipFile(path) as zf:
        slides = sorted(
            (n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[-1]).group(1)),
        )
        for idx, name in enumerate(slides, start=1):
            xml = zf.read(name).decode("utf-8", errors="replace")
            paragraphs = re.findall(r"<a:p\b.*?</a:p>", xml, re.S)
            lines = []
            for para in paragraphs:
                runs = re.findall(r"<a:t>(.*?)</a:t>", para, re.S)
                line = "".join(runs).strip()
                if line:
                    lines.append(_xml_unescape(line))
            out.pages.append({"n": idx, "text": "\n".join(lines), "is_drawing_like": False, "label": f"slide {idx}"})
    out.text = "\n\n".join(p["text"] for p in out.pages)
    out.meta["slides"] = len(out.pages)


def _xml_unescape(text: str) -> str:
    import html

    return html.unescape(text)


# --------------------------------------------------------------------------- archives
def _zip_member_name(info: zipfile.ZipInfo) -> str:
    name = info.filename
    if info.flag_bits & 0x800:
        return name
    try:
        raw = name.encode("cp437")
    except UnicodeEncodeError:
        return name
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        arabic = raw.decode("cp720")
        if re.search(r"[\u0600-\u06ff]", arabic):
            return arabic
    except UnicodeDecodeError:
        pass
    return name


def _safe_relpath(name: str) -> PurePosixPath | None:
    """Relative path inside the archive, or None for absolute / traversal / drive paths."""
    cleaned = name.replace("\\", "/")
    if cleaned.startswith("/") or re.match(r"^[A-Za-z]:", cleaned):
        return None
    parts = [p for p in cleaned.split("/") if p not in ("", ".")]
    if not parts or any(p == ".." for p in parts):
        return None
    safe = [sanitize_filename(p, "item") for p in parts]
    return PurePosixPath(*safe)


def _extract_zip(path: Path, out: DocExtract, extract_dir: Path | None, depth: int) -> None:
    target_root = _child_dir(path, extract_dir).resolve()
    members: list[dict] = []
    total = 0
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        out.warnings.append(f"not a valid ZIP archive: {exc}")
        return
    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_ZIP_MEMBERS:
            out.warnings.append(f"archive has {len(infos)} entries; only the first {MAX_ZIP_MEMBERS} were read")
            infos = infos[:MAX_ZIP_MEMBERS]
        for info in infos:
            if info.is_dir():
                continue
            name = _zip_member_name(info)
            rel = _safe_relpath(name)
            if rel is None:
                out.warnings.append(f"skipped unsafe path {name!r} (absolute or '..' path: zip-slip attempt)")
                continue
            mode = (info.external_attr >> 16) & 0o170000
            if mode == stat.S_IFLNK:
                out.warnings.append(f"skipped symbolic link {name!r}")
                continue
            if info.flag_bits & 0x1:
                out.warnings.append(f"skipped password-protected member {name!r}")
                continue
            reason = blocked_reason(rel.name)
            if reason:
                out.warnings.append(f"skipped {name!r}: {reason}")
                continue
            if info.file_size > MAX_ZIP_MEMBER_BYTES:
                out.warnings.append(f"skipped {name!r}: {info.file_size} bytes is above the member size limit")
                continue
            if info.compress_size and info.file_size > 50_000_000 and info.file_size / max(info.compress_size, 1) > 200:
                out.warnings.append(f"skipped {name!r}: suspicious compression ratio (zip bomb?)")
                continue
            if total + info.file_size > MAX_ZIP_TOTAL_BYTES:
                out.warnings.append("stopped unpacking: archive exceeds the total size limit")
                break
            destination = target_root.joinpath(*rel.parts)
            if not str(destination.resolve()).startswith(str(target_root) + os.sep):
                out.warnings.append(f"skipped unsafe path {name!r}")
                continue
            destination.parent.mkdir(parents=True, exist_ok=True)
            if destination.exists():
                destination = unique_path(destination.parent, destination.name)
            written = 0
            head = b""
            refused = None
            with archive.open(info) as src, open(destination, "wb") as dst:
                while True:
                    chunk = src.read(1 << 20)
                    if not chunk:
                        break
                    if not head:
                        head = chunk[:64]
                        refused = blocked_reason(rel.name, head)
                        if refused:
                            break
                    written += len(chunk)
                    if written > info.file_size + 1024 or written > MAX_ZIP_MEMBER_BYTES:
                        refused = "member is larger than declared (zip bomb?)"
                        break
                    dst.write(chunk)
            if refused:
                destination.unlink(missing_ok=True)
                out.warnings.append(f"skipped {name!r}: {refused}")
                continue
            total += written
            entry = {"name": rel.as_posix(), "size": written, "path": str(destination)}
            if depth < MAX_DEPTH and detect_kind(destination) not in ("unsupported",):
                child = extract_document(destination, _depth=depth + 1)
                child.name = rel.as_posix()
                out.children.append(child)
                entry.update({"kind": child.kind, "doc_kind": child.doc_kind})
                for item in child.boq_items:
                    out.boq_items.append({**item, "source": rel.as_posix()})
            else:
                entry["kind"] = detect_kind(destination)
            members.append(entry)
    out.meta.update({"members": members, "extracted_to": str(target_root), "member_count": len(members)})
    parts = [f"=== {c.name} ===\n{c.text}" for c in out.children if c.text]
    listing = "\n".join(m["name"] for m in members)
    out.text = ("Archive contents:\n" + listing + ("\n\n" + "\n\n".join(parts) if parts else "")).strip()


# --------------------------------------------------------------------------- e-mails
def _save_attachment(folder: Path, name: str, data: bytes) -> Path | None:
    folder.mkdir(parents=True, exist_ok=True)
    safe = sanitize_filename(name, "attachment")
    if blocked_reason(safe, data[:64]):
        return None
    target = unique_path(folder, safe)
    target.write_bytes(data)
    return target


def _extract_eml(path: Path, out: DocExtract, extract_dir: Path | None, depth: int) -> None:
    from ..sources.base import html_to_text, parse_address, parse_address_list

    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    body = ""
    try:
        part = msg.get_body(preferencelist=("plain", "html"))
        if part is not None:
            content = part.get_content()
            body = html_to_text(content) if part.get_content_type() == "text/html" else content
    except Exception as exc:
        out.warnings.append(f"body could not be decoded: {type(exc).__name__}")
    from_name, from_email = parse_address(str(msg.get("From") or ""))
    headers = {
        "from": from_email,
        "from_name": from_name,
        "to": parse_address_list([str(v) for v in msg.get_all("To") or []]),
        "cc": parse_address_list([str(v) for v in msg.get_all("Cc") or []]),
        "subject": str(msg.get("Subject") or ""),
        "date": str(msg.get("Date") or "") or None,
        "message_id": str(msg.get("Message-ID") or "") or None,
    }
    attachments = []
    folder = _child_dir(path, extract_dir)
    for idx, part in enumerate(msg.iter_attachments(), start=1):
        name = part.get_filename() or f"attachment-{idx}{mimetypes.guess_extension(part.get_content_type()) or '.bin'}"
        try:
            if part.get_content_type() == "message/rfc822":
                inner = part.get_content()
                data = inner.as_bytes() if hasattr(inner, "as_bytes") else bytes(inner)
                if not name.lower().endswith(".eml"):
                    name += ".eml"
            else:
                data = part.get_payload(decode=True) or b""
        except Exception as exc:
            out.warnings.append(f"attachment {name!r} could not be decoded: {type(exc).__name__}")
            continue
        saved = _save_attachment(folder, name, data)
        if saved is None:
            out.warnings.append(f"skipped attachment {name!r}: executable or script")
            continue
        attachments.append({"name": saved.name, "size": len(data), "mime": part.get_content_type(), "path": str(saved)})
        if depth < MAX_DEPTH:
            child = extract_document(saved, _depth=depth + 1)
            out.children.append(child)
            for item in child.boq_items:
                out.boq_items.append({**item, "source": saved.name})
    _finish_email(out, headers, body, attachments)


def _finish_email(out: DocExtract, headers: dict, body: str, attachments: list[dict]) -> None:
    lines = [
        f"From: {headers.get('from_name') or ''} <{headers.get('from') or ''}>".replace(" <>", ""),
        f"To: {', '.join(headers.get('to') or [])}",
    ]
    if headers.get("cc"):
        lines.append(f"Cc: {', '.join(headers['cc'])}")
    lines += [f"Date: {headers.get('date') or ''}", f"Subject: {headers.get('subject') or ''}"]
    if attachments:
        lines.append("Attachments: " + ", ".join(a["name"] for a in attachments))
    text = "\n".join(lines) + "\n\n" + (body or "").strip()
    for child in out.children:
        if child.text:
            text += f"\n\n=== attachment: {child.name} ===\n{child.text}"
    out.text = text
    out.pages = [{"n": 1, "text": "\n".join(lines) + "\n\n" + (body or "").strip(), "is_drawing_like": False}]
    out.meta.update({**headers, "attachments": attachments})


def _extract_msg(path: Path, out: DocExtract, extract_dir: Path | None, depth: int) -> None:
    import extract_msg

    from ..sources.base import html_to_text, parse_address, parse_address_list

    message = extract_msg.openMsg(str(path))
    try:
        body = message.body or ""
        if not body.strip():
            html = getattr(message, "htmlBody", None)
            if isinstance(html, bytes):
                html = html.decode("utf-8", errors="replace")
            body = html_to_text(html or "")
        from_name, from_email = parse_address(str(getattr(message, "sender", "") or ""))
        date_value = getattr(message, "date", None)
        headers = {
            "from": from_email,
            "from_name": from_name,
            "to": parse_address_list(str(getattr(message, "to", "") or "")),
            "cc": parse_address_list(str(getattr(message, "cc", "") or "")),
            "subject": str(getattr(message, "subject", "") or ""),
            "date": date_value.isoformat() if isinstance(date_value, datetime) else (str(date_value) if date_value else None),
            "message_id": getattr(message, "messageId", None),
        }
        attachments = []
        folder = _child_dir(path, extract_dir)
        for idx, att in enumerate(getattr(message, "attachments", []) or [], start=1):
            name = getattr(att, "longFilename", None) or getattr(att, "shortFilename", None) or getattr(att, "name", None) \
                or f"attachment-{idx}"
            data = getattr(att, "data", None)
            if data is not None and not isinstance(data, (bytes, bytearray)):
                exporter = getattr(data, "exportBytes", None)
                data = exporter() if callable(exporter) else None
                if data is not None and not str(name).lower().endswith(".msg"):
                    name = f"{name}.msg"
            if not data:
                continue
            saved = _save_attachment(folder, str(name), bytes(data))
            if saved is None:
                out.warnings.append(f"skipped attachment {name!r}: executable or script")
                continue
            attachments.append({"name": saved.name, "size": len(data), "path": str(saved)})
            if depth < MAX_DEPTH:
                child = extract_document(saved, _depth=depth + 1)
                out.children.append(child)
                for item in child.boq_items:
                    out.boq_items.append({**item, "source": saved.name})
    finally:
        try:
            message.close()
        except Exception:
            pass
    _finish_email(out, headers, body, attachments)


# --------------------------------------------------------------------------- public API
def extract_document(
    path: str | os.PathLike[str],
    *,
    extract_dir: str | os.PathLike[str] | None = None,
    max_pages: int = 2000,
    _depth: int = 0,
) -> DocExtract:
    """Extract one file. Raises ``FileNotFoundError`` for a missing path; everything else
    (corrupt, encrypted, unsupported) is reported in ``warnings``."""
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(str(file_path))
    kind = detect_kind(file_path)
    out = DocExtract(path=str(file_path), name=file_path.name, kind=kind)
    out.meta["size"] = file_path.stat().st_size
    child_dir = Path(extract_dir) if extract_dir is not None else None
    try:
        if kind == "pdf":
            parts = extract_pdf(file_path, max_pages=max_pages)
            for key, value in parts.items():
                if key == "meta":
                    out.meta.update(value)
                elif key == "warnings":
                    out.warnings += value
                else:
                    setattr(out, key, value)
        elif kind == "docx":
            _extract_docx(file_path, out)
        elif kind == "spreadsheet":
            ext = file_path.suffix.lower()
            if ext == ".xls":
                _extract_xls(file_path, out)
            elif ext in (".xlsb", ".ods"):
                out.warnings.append(f"{ext} workbooks are not supported; save as .xlsx")
            else:
                _extract_xlsx(file_path, out)
        elif kind == "csv":
            _extract_csv(file_path, out)
        elif kind == "image":
            _extract_image(file_path, out)
        elif kind == "archive":
            if file_path.suffix.lower() in (".zip", "") or zipfile.is_zipfile(file_path):
                _extract_zip(file_path, out, child_dir, _depth)
            else:
                out.warnings.append(f"{file_path.suffix} archives are not supported; unpack it and add the files")
        elif kind == "email":
            if file_path.suffix.lower() == ".msg":
                _extract_msg(file_path, out, child_dir, _depth)
            else:
                _extract_eml(file_path, out, child_dir, _depth)
        elif kind == "cad":
            parts = read_cad(file_path)
            out.text = parts["text"]
            out.meta.update(parts["meta"])
            out.warnings += parts["warnings"]
        elif kind == "text":
            _extract_text_file(file_path, out)
        elif kind == "presentation":
            _extract_pptx(file_path, out)
        else:
            hint = {".doc": "legacy .doc: save as .docx or PDF", ".ppt": "legacy .ppt: save as .pptx or PDF"}
            out.warnings.append(hint.get(file_path.suffix.lower(), f"unsupported file type {file_path.suffix or '(none)'}"))
    except Exception as exc:  # corrupt or unexpected content: report, don't crash the pipeline
        out.warnings.append(f"extraction failed: {type(exc).__name__}: {exc}")
    if len(out.text) > MAX_TEXT_CHARS:
        out.text = out.text[:MAX_TEXT_CHARS]
        out.warnings.append(f"text truncated to {MAX_TEXT_CHARS} characters")
    out.meta.setdefault("pages", len(out.pages))
    out.meta["boq_relevant"] = sum(1 for b in out.boq_items if b.get("relevant"))
    out.doc_kind = classify_document(file_path.name, out.text, boq_items=out.boq_items, pages=out.pages)
    return out


# --------------------------------------------------------------------------- classification
_CAD_EXT = {".dwg", ".dxf", ".rvt", ".rfa", ".ifc", ".dgn", ".skp", ".nwd", ".nwc", ".dwf", ".dwfx", ".step", ".stp"}
_PHOTO_EXT = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
_DRAWING_NAME_RE = re.compile(
    r"(\b(dwg|drg|drawings?|layout|elevations?|sections?|plans?|details?|shop\s*drawings?|sheet)\b"
    # sheet numbers such as A-101, AR-201, S-1001, M 301 (not years, not RFQ-/BOQ-/PO- references)
    r"|(^|[\s_(\-])(?!(rfq|rfp|boq|sor|ref|qtn|quo|inv|lpo|po|no|itt|tdr|img|dsc|pxl)[-_ ]?\d)"
    r"([A-Z]{1,3}|\d{2})[-_ ]?(?!(19|20)\d{2}\b)\d{3,4}([-_ ]?(rev|r)\.?\s?[A-Z0-9]{1,2})?\b)",
    re.I,
)
_BOQ_NAME_RE = re.compile(
    r"(\bboq\b|\bb\.o\.q\b|bills?\s*of\s*quantit|pricing\s*schedule|price\s*schedule|schedule\s*of\s*rates|\bsor\b"
    r"|priced\s*bill|جدول\s*الكميات|الكميات)",
    re.I,
)
_SPEC_NAME_RE = re.compile(r"(spec|specification|specs\b|technical\s*requirements|section\s*\d{2}\s?\d{2}|مواصفات|اشتراطات)", re.I)
_ADDENDUM_RE = re.compile(r"(addend(um|a)|corrigendum|clarification|circular|bulletin|ملحق|تعميم|توضيح)", re.I)
_TENDER_RE = re.compile(
    r"(\btender\b|\brfq\b|\brfp\b|\bitt\b|invitation\s*to\s*(bid|tender)|instructions\s*to\s*(bidders|tenderers)"
    r"|conditions\s*of\s*contract|form\s*of\s*tender|request\s*for\s*(quotation|proposal)|مناقص|كراس[ةه]\s*الشروط|عطاء)",
    re.I,
)
_SPEC_TEXT_RE = re.compile(
    r"(part\s*1\s*[-–.]?\s*general|section\s*\d{2}\s?\d{2}\s?\d{2}|1\.\d+\s+(summary|scope|references|submittals)"
    r"|quality\s*assurance|submittals|performance\s*requirements|execution\s*$)",
    re.I | re.M,
)
_LETTER_RE = re.compile(r"(dear\s+(sir|madam|mr|ms|all)|yours\s+(faithfully|sincerely)|kind\s+regards|best\s+regards|السادة|تحية\s*طيبة)", re.I)
_TITLE_WORDS_RE = re.compile(r"\b(drawing\s*no|dwg\s*no|drg\s*no|scale|checked|drawn|revision|rev)\b", re.I)


def classify_document(name: str, text: str = "", *, boq_items: list[dict] | None = None,
                      pages: list[dict] | None = None) -> str:
    """drawing | boq | specification | tender_doc | addendum | photo | cad | correspondence | other."""
    name = name or ""
    stem = Path(name).stem
    ext = Path(name).suffix.lower()
    head = (text or "")[:6000]
    if ext in _CAD_EXT:
        return "cad"
    if ext in (".eml", ".msg"):
        return "correspondence"
    if _ADDENDUM_RE.search(stem) or re.search(r"^\s*(addendum|corrigendum)\s*(no\.?|#)?\s*\d*", head, re.I | re.M):
        return "addendum"
    if ext in _PHOTO_EXT:
        return "drawing" if _DRAWING_NAME_RE.search(stem) and not re.search(r"(img|dsc|photo|whatsapp|pxl)", stem, re.I) else "photo"
    if _BOQ_NAME_RE.search(stem):
        return "boq"
    if _SPEC_NAME_RE.search(stem):
        return "specification"
    if re.search(r"\b(dwg|drg|drawings?|shop\s*drawings?|layouts?|elevations?)\b", stem, re.I):
        return "drawing"
    if _TENDER_RE.search(stem):
        return "tender_doc"
    drawing_pages = [p for p in (pages or []) if p.get("is_drawing_like")]
    if pages and len(drawing_pages) * 2 > len(pages):
        return "drawing"
    if _DRAWING_NAME_RE.search(stem) and (not text or len(_TITLE_WORDS_RE.findall(head)) >= 2 or drawing_pages):
        return "drawing"
    if boq_items and len(boq_items) >= 3:
        return "boq"
    if re.search(r"bills?\s+of\s+quantities|schedule\s+of\s+rates|pricing\s+schedule", head, re.I):
        return "boq"
    if _SPEC_TEXT_RE.search(head) and len(_SPEC_TEXT_RE.findall(head)) >= 2:
        return "specification"
    if len(_TENDER_RE.findall(head)) >= 2:
        return "tender_doc"
    if len(_TITLE_WORDS_RE.findall(head)) >= 3 and len(head) < 3000:
        return "drawing"
    if _LETTER_RE.search(head):
        return "correspondence"
    if is_relevant(stem) and re.search(r"(spec|requirement)", head, re.I):
        return "specification"
    return "other"
