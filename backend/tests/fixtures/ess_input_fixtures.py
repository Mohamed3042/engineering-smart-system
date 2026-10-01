"""Generators for the input-side test fixtures (documents are built at test time, never committed)."""
from __future__ import annotations

import io
import zipfile
from email.message import EmailMessage
from pathlib import Path

MM = 72.0 / 25.4  # points per millimetre


# --------------------------------------------------------------------------- PDF
def make_pdf(path: Path, pages: list[dict], *, title: str | None = None, author: str | None = None) -> Path:
    """Text PDF from ``[{"width": pt, "height": pt, "lines": [(x, y, size, text), ...]}]`` (Helvetica)."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
            NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
        }
    )
    font_ref = writer._add_object(font)
    for spec in pages:
        page = writer.add_blank_page(width=spec["width"], height=spec["height"])
        ops = []
        for x, y, size, text in spec["lines"]:
            esc = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"BT /F1 {size} Tf {x:.1f} {y:.1f} Td ({esc}) Tj ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(ops).encode("cp1252"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
    meta = {"/CreationDate": "D:20260826082026Z"}
    if title:
        meta["/Title"] = title
    if author:
        meta["/Author"] = author
    writer.add_metadata(meta)
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def spec_page() -> dict:
    """A4 portrait specification page: dense text."""
    lines = [(72, 790, 12, "SECTION 11 24 23 - FACADE ACCESS EQUIPMENT"), (72, 770, 10, "PART 1 - GENERAL"),
             (72, 755, 10, "1.1 SUMMARY")]
    body = (
        "A. This Section includes roof-mounted building maintenance units, monorail systems, davit arms and "
        "suspended cradles for facade cleaning and maintenance of the tower. "
    )
    y = 738
    for i in range(44):
        lines.append((72, y, 8, f"{i + 1:02d} {body[(i * 7) % 60:][:95]}"))
        y -= 15
    return {"width": 210 * MM, "height": 297 * MM, "lines": lines}


def drawing_page() -> dict:
    """A1 landscape drawing sheet with a title block and sparse annotations."""
    w, h = 841 * MM, 594 * MM
    tb_x = w - 420
    lines = [
        (300, 1200, 14, "ROOF PLAN - BMU TRACK AND GONDOLA PARKING"),
        (420, 900, 10, "BMU RAIL"),
        (900, 640, 10, "PARKING BAY"),
        (600, 400, 10, "DAVIT SOCKET (TYP.)"),
        (tb_x, 230, 9, "PROJECT: HARBOUR TOWER, PLOT 12"),
        (tb_x, 210, 9, "CLIENT: EXAMPLE CONTRACTING CO."),
        (tb_x, 190, 9, "DRAWING TITLE: ROOF BMU LAYOUT"),
        (tb_x, 170, 9, "DRAWING NO: A-501"),
        (tb_x, 150, 9, "SCALE: 1:100 @ A1"),
        (tb_x, 130, 9, "REV: B"),
        (tb_x, 110, 9, "DRAWN: MK   CHECKED: AI   APPROVED: HS"),
        (tb_x, 90, 9, "ISSUED FOR TENDER"),
    ]
    return {"width": w, "height": h, "lines": lines}


def boq_page() -> dict:
    """A3 landscape BOQ page with column-aligned text."""
    w, h = 420 * MM, 297 * MM
    cols = {"item": 40, "desc": 110, "qty": 700, "unit": 780, "rate": 860, "amount": 980}
    lines = [(40, 800, 12, "BILL OF QUANTITIES - HARBOUR TOWER")]
    header = [("item", "Item"), ("desc", "Description"), ("qty", "Qty"), ("unit", "Unit"), ("rate", "Rate"),
              ("amount", "Amount")]
    for key, label in header:
        lines.append((cols[key], 760, 10, label))
    rows = [
        {"desc": "DIVISION 11 - EQUIPMENT"},
        {"item": "11 24 23", "desc": "Facade Access Equipment"},
        {"item": "11.1", "desc": "Building Maintenance Unit (BMU), roof-mounted, telescopic jib", "qty": "1", "unit": "No"},
        {"item": "11.2", "desc": "Monorail system for window cleaning, stainless steel", "qty": "120", "unit": "m"},
        {"desc": "DIVISION 08 - OPENINGS"},
        {"item": "8.1", "desc": "Aluminium entrance doors", "qty": "12", "unit": "No", "rate": "350.000",
         "amount": "4,200.000"},
    ]
    y = 735
    for row in rows:
        for key, value in row.items():
            lines.append((cols[key], y, 9, value))
        y -= 22
    return {"width": w, "height": h, "lines": lines}


def make_tender_pdf(path: Path) -> Path:
    return make_pdf(path, [spec_page(), drawing_page(), boq_page()], title="Tender package", author="UE Estimation")


# --------------------------------------------------------------------------- Office
def make_docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    document.core_properties.author = "A. Engineer"
    document.core_properties.title = "RFQ - BMU"
    document.add_heading("Request for Quotation", level=1)
    document.add_paragraph("Please quote for the supply and installation of a building maintenance unit (BMU) "
                           "and a window cleaning monorail for Harbour Tower.")
    table = document.add_table(rows=1, cols=6)
    for cell, label in zip(table.rows[0].cells, ["Item", "Description", "Qty", "Unit", "Rate", "Amount"]):
        cell.text = label
    for values in (["1", "Roof-mounted BMU with cradle", "1", "No", "", ""],
                   ["2", "Painting of plant room", "1", "LS", "", ""]):
        row = table.add_row().cells
        for cell, value in zip(row, values):
            cell.text = value
    section = document.sections[0]
    section.header.paragraphs[0].text = "UE/RFQ/2026/118"
    document.save(str(path))
    return path


def make_boq_xlsx(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Div 11 - Equipment"
    rows = [
        ["HARBOUR TOWER - BILL OF QUANTITIES"],
        ["Tender No. UE/2026/118"],
        [],
        ["Item", "Description", "Qty", "Unit", "Rate (KD)", "Amount (KD)"],
        [None, "DIVISION 11 – EQUIPMENT"],
        ["11 24 23", "Façade Access Equipment"],
        ["11.1", "Supply, install, test and commission roof-mounted Building Maintenance Unit (BMU) "
                 "with telescopic jib", 1, "No", None, None],
        [None, "complete with cradle, control panel and all accessories as per Section 11 24 23."],
        ["11.2", "Monorail system for window cleaning of atrium glazing", 45, "m", None, None],
        ["11.3", "Davit arms with sockets", 6, "Nos", None, None],
        ["11 13 00", "Loading Dock Equipment"],
        ["11.4", "Dock leveller, hydraulic", 2, "No", 1500, 3000],
        [None, "Total carried to collection", None, None, None, 3000],
        ["DIVISION 08 – OPENINGS"],
        ["8.1", "Aluminium entrance doors", 4, "No", None, None],
    ]
    for row in rows:
        ws.append(row)
    ar = wb.create_sheet("ملخص")
    ar.append(["البند", "البيان", "الكمية", "الوحدة", "سعر الوحدة", "المبلغ"])
    ar.append(["1", "وحدة صيانة المباني (BMU) مع الملحقات", 1, "عدد", None, None])
    ar.append(["2", "أعمال الدهانات", 100, "م2", 2.5, 250])
    wb.save(str(path))
    return path


# --------------------------------------------------------------------------- archives & mail
def make_zip_with_slip(path: Path, members: dict[str, bytes]) -> Path:
    """ZIP with the given members plus zip-slip / absolute-path / executable entries."""
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
        zf.writestr(zipfile.ZipInfo("../../evil.txt"), b"zip slip")
        zf.writestr(zipfile.ZipInfo("/tmp/absolute-evil.txt"), b"absolute")
        zf.writestr(zipfile.ZipInfo("tools/setup.exe"), b"MZ\x90\x00 fake exe")
        zf.writestr(zipfile.ZipInfo("tools/disguised.pdf"), b"MZ\x90\x00 really an exe")
    return path


def make_eml(path: Path, attachment_name: str, attachment: bytes, mime: str) -> Path:
    msg = EmailMessage()
    msg["From"] = "A. Engineer <a.engineer@contractor.example>"
    msg["To"] = "sales@medmack.com"
    msg["Subject"] = "RFQ – BMU (طلب تسعير)"
    msg["Date"] = "Wed, 26 Aug 2026 11:20:26 +0300"
    msg["Message-ID"] = "<eml-test-1@contractor.example>"
    msg.set_content("Dear Sir,\nPlease find the BOQ attached for the BMU works.\nRegards")
    maintype, subtype = mime.split("/", 1)
    msg.add_attachment(attachment, maintype=maintype, subtype=subtype, filename=attachment_name)
    path.write_bytes(bytes(msg))
    return path


def make_dxf(path: Path) -> Path:
    pairs = [
        ("0", "SECTION"), ("2", "HEADER"), ("9", "$ACADVER"), ("1", "AC1027"), ("0", "ENDSEC"),
        ("0", "SECTION"), ("2", "ENTITIES"),
        ("0", "TEXT"), ("8", "A-ANNO"), ("10", "0.0"), ("20", "0.0"), ("40", "2.5"), ("1", "ROOF BMU LAYOUT"),
        ("0", "MTEXT"), ("8", "A-ANNO"), ("3", "{\\fArial|b1|i0;BMU TRACK}\\P"), ("1", "DETAIL A %%c50"),
        ("0", "ATTRIB"), ("8", "TITLE"), ("2", "DWG_NO"), ("1", "A-501"),
        ("0", "DIMENSION"), ("1", "<>"),
        ("0", "LINE"), ("8", "0"), ("10", "0"), ("20", "0"), ("11", "10"), ("21", "10"),
        ("0", "TEXT"), ("1", "\\U+0645\\U+0642\\U+064A\\U+0627\\U+0633 1:100"),
        ("0", "ENDSEC"), ("0", "EOF"),
    ]
    path.write_text("\n".join(f"{code}\n{value}" for code, value in pairs) + "\n", encoding="utf-8")
    return path


def zip_bytes(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()
