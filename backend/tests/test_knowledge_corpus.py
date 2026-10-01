"""ess.knowledge.corpus: corpus documents from folders and mailbox messages."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import pytest

from ess.knowledge import corpus
from ess.knowledge.corpus import SOURCE_WEIGHTS, CorpusDoc, build_corpus_from_folder, from_mail_messages

FIXTURES = Path(__file__).parent / "fixtures" / "knowledge"


def _make_pdf(text: str) -> bytes:
    """A minimal valid one-page PDF with a line of Helvetica text."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


@pytest.fixture(autouse=True)
def no_documents_package(monkeypatch):
    """Use the built-in readers unless a test installs a fake ess.documents extractor."""
    monkeypatch.setattr(corpus, "_documents_extractor", lambda: None)


def test_corpus_doc_basics():
    d = CorpusDoc(source_type="own_quotation", source_id="q1", label="Q1", text="Supply of one BMU",
                  date=datetime(2026, 8, 1, tzinfo=timezone.utc))
    assert d.weight == SOURCE_WEIGHTS["own_quotation"] == 1.0 and d.is_own
    assert d.date.startswith("2026-08-01") and d.language == "en"
    assert SOURCE_WEIGHTS == {"own_quotation": 1.0, "company_doc": 0.9, "sent_email": 0.8, "inbound_email": 0.5,
                              "web": 0.2}
    with pytest.raises(ValueError):
        CorpusDoc(source_type="rumour", source_id="x", label="x", text="x")


def test_build_corpus_from_fixture_folder():
    docs = build_corpus_from_folder(FIXTURES / "northwind" / "quotations", source_type="own_quotation")
    assert len(docs) == 3
    assert all(d.source_type == "own_quotation" and d.text.startswith("NORTHWIND") for d in docs)
    assert docs[0].source_id == "NW-26-0101 BMU Al Noor Tower.txt" and docs[0].label.endswith(".txt")
    assert len(build_corpus_from_folder(FIXTURES / "northwind", limit=2)) == 2


def test_build_corpus_reads_docx_pdf_html_and_auto_type(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("QUOTATION Ref: AA/26/0118")
    d.add_paragraph("We are pleased to offer the supply of davit systems.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Davit arm"
    table.rows[0].cells[1].text = "4 nos"
    d.save(tmp_path / "offer.docx")
    (tmp_path / "manual.pdf").write_bytes(_make_pdf("Installation manual for monorail systems"))
    (tmp_path / "about.html").write_text("<html><body><nav>Home</nav><p>We maintain BMUs.</p>"
                                         "<script>x=1</script></body></html>", encoding="utf-8")
    (tmp_path / "photo.jpg").write_bytes(b"\xff\xd8not a document")
    (tmp_path / ".hidden.txt").write_text("secret", encoding="utf-8")
    docs = {d.source_id: d for d in build_corpus_from_folder(tmp_path, source_type="auto")}
    assert set(docs) == {"offer.docx", "manual.pdf", "about.html"}
    assert docs["offer.docx"].source_type == "own_quotation"
    assert "Davit arm | 4 nos" in docs["offer.docx"].text
    assert "monorail systems" in docs["manual.pdf"].text and docs["manual.pdf"].source_type == "company_doc"
    assert "We maintain BMUs." in docs["about.html"].text and "x=1" not in docs["about.html"].text


def test_build_corpus_prefers_documents_extractor(monkeypatch, tmp_path):
    (tmp_path / "boq.xlsx").write_bytes(b"PK fake")

    class Extract:
        text = "Item 1 Supply of cradle"

    monkeypatch.setattr(corpus, "_documents_extractor", lambda: (lambda path: Extract()))
    docs = build_corpus_from_folder(tmp_path)
    assert [d.text for d in docs] == ["Item 1 Supply of cradle"]


def test_build_corpus_missing_folder():
    with pytest.raises(FileNotFoundError):
        build_corpus_from_folder("/nonexistent/folder/for/ess")


def test_from_mail_messages_splits_directions_and_strips_history():
    mail = json.loads((FIXTURES / "northwind" / "mail.json").read_text(encoding="utf-8"))
    docs = {d.source_id: d for d in from_mail_messages(mail["messages"], mail["own_domains"])}
    assert docs["s1"].source_type == "sent_email" and docs["s2"].source_type == "sent_email"  # SENT label
    assert docs["s3"].source_type == "sent_email"
    assert "s4" not in docs  # forward without own words
    assert "gondola system" not in docs["s1"].text  # the customer's quoted request is not our wording
    assert docs["s1"].text.startswith("Quotation NW/26/0101")  # new outbound subject kept
    assert "Cherry picker hire extension" not in docs["s2"].text.splitlines()[0]  # reply subject dropped
    assert "man lift rental" not in docs["s2"].text  # quoted Outlook history cut
    assert docs["i1"].source_type == "inbound_email" and docs["i1"].meta["from_domain"] == "alsafwa-kw.example"
    assert docs["i5"].text == "Thank you, noted. We will revert after the client meeting."
    assert docs["i3"].language == "ar"


def test_from_mail_messages_accepts_mail_message_models():
    base = pytest.importorskip("ess.sources.base")
    msgs = [
        base.MailMessage(id="m1", thread_id="t", from_email="me@acme.example", subject="Offer for hoists",
                         date=datetime(2026, 9, 1, tzinfo=timezone.utc), body_text="We offer two hoists."),
        base.MailMessage(id="m2", thread_id="t", from_email="buyer@client.example", subject="RFQ hoists",
                         date=datetime(2026, 9, 2, tzinfo=timezone.utc), body_text="Please quote two hoists."),
    ]
    docs = from_mail_messages(msgs, ["acme.example"])
    assert [(d.source_id, d.source_type) for d in docs] == [("m1", "sent_email"), ("m2", "inbound_email")]
    assert docs[0].date.startswith("2026-09-01")
